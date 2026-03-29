import streamlit as st
import pdfplumber
import os
import streamlit as st
import pdfplumber
import os
import re
from openai import OpenAI
from dotenv import load_dotenv
from docx import Document
from io import BytesIO

# 1. Setup & API Key
# override=True ensures it reloads the key if you change your .env file
load_dotenv(override=True)
client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

# 2. Function to read PDF Resume
def extract_text_from_pdf(file):
    try:
        with pdfplumber.open(file) as pdf:
            text = ""
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        return text
    except Exception as e:
        return f"Error reading PDF: {e}"

# 3. Function to Create Word Document (.docx)
def create_word_doc(analysis_text):
    doc = Document()
    doc.add_heading('Group 3: Resume Optimization Report', 0)
    
    # Remove markdown bolding (**) for a cleaner Word look
    clean_text = analysis_text.replace("**", "")
    
    # Add the AI content to the document
    doc.add_paragraph(clean_text)
    
    # Save to a memory buffer (BytesIO) so it can be downloaded
    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()

# 4. UI Design (The Frontend)
st.set_page_config(page_title="Career Optimizer", page_icon="🚀", layout="wide")

# Header Section
st.title("Group 3: Resume & Career Optimizer")
st.markdown("#### Generative AI-Powered Career Assistant")
st.info("Upload your resume and a job description to identify skill gaps and optimize your profile.")

# Layout: Two columns for inputs
col_input1, col_input2 = st.columns(2)

with col_input1:
    st.header("1. Upload Resume")
    uploaded_file = st.file_uploader("Choose a PDF file", type="pdf")

with col_input2:
    st.header("2. Job Description")
    jd_text = st.text_area("Paste the target Job Description here", height=200)

# 5. Main Analysis Logic
if st.button("Analyze & Optimize"):
    if uploaded_file and jd_text:
        with st.spinner('AI is analyzing requirements and mapping your skills...'):
            # Extract text from the uploaded PDF
            resume_content = extract_text_from_pdf(uploaded_file)
            
            # THE AI PROMPT: Designed to meet all assignment requirements
            prompt = f"""
            You are an expert Executive Career Coach and Technical Recruiter. 
            Compare the provided Resume against the Job Description. 

            Please provide the following four sections in your response:

            1. **Match Score**: Provide a match percentage (0-100%) and a brief explanation why.
            2. **Skill Gap Analysis**: List specific technical and soft skills mentioned in the Job Description that are missing or weak in the Resume.
            3. **Suggested Improvements**: Provide 3-5 actionable tips to make the resume more competitive for this specific role.
            4. **Professional Bullet Rewrites**: Identify 3 actual bullet points from the resume and rewrite them using the 'X-Y-Z' formula (Accomplished [X] as measured by [Y], by doing [Z]). Make them quantified and impactful.

            Resume Content:
            {resume_content}

            Job Description:
            {jd_text}
            """
            
            try:
                # Call OpenAI API
                response = client.chat.completions.create(
                    model="gpt-4o-mini",
                    messages=[
                        {"role": "system", "content": "You are a professional career optimizer assistant."},
                        {"role": "user", "content": prompt}
                    ],
                    temperature=0.7
                )
                
                analysis_result = response.choices[0].message.content

                # --- DISPLAY RESULTS ---
                st.success("Analysis Complete!")
                st.divider()

                # Visualizing the Match Score using a metric
                score_match = re.search(r"(\d+)%", analysis_result)
                score = score_match.group(1) if score_match else "N/A"
                
                st.metric(label="Overall Match Score", value=f"{score}%")
                
                # Show the full AI feedback in the app
                st.markdown(analysis_result)
                
                # --- WORD DOWNLOAD LOGIC ---
                # Generate the Word doc in memory
                word_data = create_word_doc(analysis_result)
                
                # Updated Download Button for Word
                st.download_button(
                    label="📥 Download Optimization Report (Word)",
                    data=word_data,
                    file_name="Resume_Optimization_Report.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )

            except Exception as e:
                st.error(f"An error occurred with the AI service: {e}")
                
    else:
        st.warning("⚠️ Please provide both a Resume (PDF) and a Job Description to proceed.")

# 6. Footer (Course Info)
st.divider()
st.caption("Executive Master’s in Artificial Intelligence | University of the Cumberlands | Group 3 Project")